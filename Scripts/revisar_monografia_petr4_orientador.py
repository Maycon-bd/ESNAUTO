"""Prepara a revisão solicitada pelo orientador sem alterar as versões anteriores."""
from pathlib import Path
from copy import deepcopy
import sys,json,hashlib
import pandas as pd
import numpy as np
sys.path.append(str(Path(__file__).resolve().parents[1]/'reports/vale3_comparacao/python_deps'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/revisao_orientador_petr4'
SOURCE=ROOT/'reports/monografia/MayconGarciaSilva_monografia_bibliometria_VALE3.docx'
DEST=ROOT/'reports/monografia/MayconGarciaSilva_monografia_PETR4_revisao_orientador.docx'
source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
h=pd.read_csv(OUT/'historico_ga.csv');p=pd.read_csv(OUT/'previsoes_teste.csv',parse_dates=['data'])
m=pd.read_csv(OUT/'metricas.csv').iloc[0]
assert len(h)==200 and len(p)==1298
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11})
fig,ax=plt.subplots(figsize=(9,3.6),layout='constrained')
ax.plot(h.geracao,h.fitness_minimo,label='Mínimo da população',color='#999999',lw=.7,alpha=.65)
ax.plot(h.geracao,h.fitness_medio,label='Média da população',color='#28769e',lw=1)
ax.step(h.geracao,h.melhor_fitness,label='Melhor fitness da geração',color='#188451',lw=2,where='post')
ax.axvline(123,color='#874e2f',ls='--',lw=1)
ax.annotate('Melhor valor atingido\nna geração 123',xy=(123,h.melhor_fitness.max()),xytext=(137,-.23768),fontsize=9,arrowprops={'arrowstyle':'->','color':'#874e2f'})
ax.set(xlabel='Geração',ylabel='Fitness = −(0,4 MAE treino + 0,6 MAE validação)',xlim=(1,200))
ax.ticklabel_format(axis='y',style='plain',useOffset=False);ax.grid(alpha=.2);ax.legend(loc='lower right',fontsize=9)
fig.savefig(OUT/'fig5_evolucao_ga_petr4.png',dpi=230);fig.savefig(OUT/'fig5_evolucao_ga_petr4.pdf');plt.close(fig)
fig,axs=plt.subplots(2,1,figsize=(9,5),layout='constrained')
for ax,data,title in [(axs[0],p,'Conjunto de teste completo'),(axs[1],p.tail(100),'Detalhe dos últimos 100 pregões avaliados')]:
    ax.plot(data.data,data.observado,color='#242424',lw=1.4,label='Observado')
    ax.plot(data.data,data.previsto,color='#297daa',lw=1,ls='--',label='Previsto')
    ax.set_title(title,fontsize=11);ax.set_ylabel('Preço ajustado (R$)');ax.grid(alpha=.2)
    ax.set_xlim(data.data.min(),data.data.max())
    if ax is axs[0]:
        ax.xaxis.set_major_locator(mdates.YearLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%Y'))
    else:
        ax.xaxis.set_major_locator(mdates.MonthLocator())
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%m/%Y'))
axs[0].legend(ncol=2,fontsize=10);axs[1].set_xlabel('Data')
fig.savefig(OUT/'fig6_observado_previsto_petr4.png',dpi=230);fig.savefig(OUT/'fig6_observado_previsto_petr4.pdf');plt.close(fig)

d=Document(SOURCE);original=list(d.paragraphs)
def replace(p,text):
    prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear();r=p.add_run(text)
    if prop is not None:r._r.insert(0,prop)
def remove(p):p._element.getparent().remove(p._element)
for i in [215,399,400,525,539]:remove(original[i])
remove(original[462])  # Parágrafo vazio que, com a quebra do capítulo, gerava uma página em branco.
for p0 in original:
    if 'Os cenários reduzidos de 12 combinações e 200 gerações pertenceram' in p0.text:
        replace(p0,p0.text.replace('Os cenários reduzidos de 12 combinações e 200 gerações pertenceram à etapa de prototipação e não integraram os resultados finais.',
            'O lote reduzido de 12 cenários e 200 gerações pertenceu à prototipação e não integra as 45 execuções de produção; seus registros preservados são utilizados somente nas ilustrações complementares das Figuras 5 e 6.'))
replace(original[545],'Como continuidade experimental e possível desdobramento em artigo, propõe-se ampliar o estudo à ação VALE3, reproduzindo a caracterização da série, o pré-processamento, as combinações de distribuições, as repetições e a comparação com LSTM e GRU realizadas para PETR4. As execuções exploratórias já preservadas do segundo ativo constituem material preliminar e não integram os resultados desta monografia. Essa ampliação requer partições temporalmente auditáveis e parâmetros das distribuições ajustados apenas ao treino, ou identificação explícita de sua transferência entre ativos.')
replace(original[551],original[551].text.split(' A extensão do seletor para PETR4 e VALE3')[0])
replace(original[461],original[461].text.replace(
 'Na execução complementar destinada ao teste de Diebold-Mariano, essa limitação foi tratada prospectivamente: semente, versões de dependências, matrizes, modelos, previsões datadas, métricas, validações e hashes foram persistidos antes da análise inferencial.',
 'A execução complementar de Diebold-Mariano preservou semente, dependências, matrizes, modelos, previsões datadas, métricas, validações e hashes antes da análise inferencial.'))
anchor=original[478]
def before(text='',style='Normal'):
    el=OxmlElement('w:p');anchor._p.addprevious(el);q=Paragraph(el,anchor._parent);q.style=style
    if style=='Normal' and original[477]._p.pPr is not None:q._p.insert(0,deepcopy(original[477]._p.pPr))
    q.add_run(text)
    for r in q.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
    q.paragraph_format.widow_control=True
    return q
texts=[
 'Para complementar as tabelas, as Figuras 5 e 6 apresentam a evolução da busca genética e a comparação entre valores observados e previstos na mesma execução preservada. A escolha dessas visualizações tomou como referência as Figuras 29 e 30 de Cardoso (2025, p. 90), adaptando-as à previsão do preço ajustado, em vez dos retornos estudados naquela tese.',
 'A execução ilustrada pertence ao lote legado de 12 cenários com limite de 200 gerações, de 15 de agosto de 2026. Foi selecionada pelo maior fitness de treino/validação nesse lote: segunda rodada de GED em Win e Uniforme em W, com a = 0,424701, raio espectral = 0,513577, lavagem de 34 passos, 31 neurônios e regularização de aproximadamente 4,84 × 10⁻⁵. Trata-se de uma ilustração complementar do motor legado, distinta das 45 execuções de produção com LHS e cataclismo. Não constitui a reconstrução da trajetória do GA_RUN_0030 nem substitui os resultados da Tabela 4.',
 'Na Figura 5, valores maiores de fitness correspondem a menores erros ponderados de treino e validação. O melhor fitness atingiu −0,237607 na geração 123 e permaneceu nesse nível até a geração 200. As oscilações das curvas de média e mínimo descrevem a variação da população durante a busca; a estabilização do melhor valor não demonstra que tenha sido encontrado um ótimo global.',
 'A Figura 6 foi produzida pela reinjeção das matrizes Win, W e Wout preservadas da mesma execução, sem gerar novos pesos aleatórios nem repetir o GA. As previsões foram alinhadas aos 1.298 alvos avaliados do conjunto de teste e as métricas foram conferidas contra o relatório arquivado: MAE de 0,327894 e RMSE de 0,498699. A apresentação inclui o período completo e um detalhe fixo dos últimos 100 pregões, permitindo observar desvios que ficam pouco visíveis na escala de toda a série.',
 'A proximidade visual entre as curvas deve ser interpretada em conjunto com os erros numéricos: o gráfico descreve a previsão de um passo à frente em preços em nível e não representa uma trajetória inteiramente recursiva de longo prazo. O teste não participou da escolha da execução ilustrada. Como os históricos completos de evolução do GA das execuções de produção não foram preservados, não se atribui a elas a curva de convergência deste lote legado.'
]
before(texts[0]);before(texts[1])
caps=['Figura 5 – Evolução do fitness na execução preservada da PETR4 com 200 gerações',
      'Figura 6 – Preços observados e previstos da PETR4 na mesma execução preservada']
def figure(caption,path,source):
    q=before(caption,'Caption');q.paragraph_format.keep_with_next=True
    q=before();q.paragraph_format.first_line_indent=Cm(0);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.add_run().add_picture(str(path),width=Cm(15));q.paragraph_format.keep_with_next=True
    q=before(source);q.paragraph_format.first_line_indent=Cm(0)
    for r in q.runs:r.font.size=Pt(10)
figure(caps[0],OUT/'fig5_evolucao_ga_petr4.png','Fonte: Elaborado pelo autor com o histórico preservado do lote de 200 gerações (2026). Referência de apresentação: Cardoso (2025, Figura 29).')
before(texts[2]);before(texts[3])
figure(caps[1],OUT/'fig6_observado_previsto_petr4.png','Fonte: Elaborado pelo autor pela reinjeção das matrizes preservadas, GED × Uniforme, rodada 2 (2026). Referência de apresentação: Cardoso (2025, Figura 30).')
before(texts[4])
# Delimita a origem desse material na metodologia e na discussão de rastreabilidade.
el=OxmlElement('w:p');original[434]._p.addnext(el);q=Paragraph(el,original[434]._parent);q.style='Normal'
q.add_run('Como material visual complementar, utilizou-se também um lote legado de 12 cenários com 200 gerações, cujo histórico de fitness e matrizes foram preservados. A escolha para as Figuras 5 e 6 considerou exclusivamente o maior fitness de treino/validação desse lote. Esses registros não foram agregados às 45 execuções de produção, e a ilustração não replica o protocolo com LHS e cataclismo.')
for r in q.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
# Lista de figuras estática: mantém o desenho da versão revisada anteriormente.
last=next(p for p in original if p.text.startswith('Figura 4') and '\t' in p.text)
for caption in caps:
    el=deepcopy(last._p);last._p.addnext(el);q=Paragraph(el,last._parent);replace(q,caption+'\t0');last=q
# A tese é referência de apresentação; não altera a coleção bibliométrica de 34 documentos.
ref=original[562]
el=deepcopy(ref._p);ref._p.addnext(el);q=Paragraph(el,ref._parent)
replace(q,'CARDOSO, F. C. Modelo híbrido para previsão de séries temporais financeiras: ARIMA-ESN com algoritmo genético. 2025. Tese (Doutorado em Modelagem Computacional) – Universidade Federal do Rio Grande, Rio Grande, 2025. Disponível em: https://ppgmc.furg.br/dissertacoes-e-teses/publicacoes-de-2025/14475tese-fabian-correa-cardoso. Acesso em: 23 out. 2026.')
d.save(DEST)
check=Document(DEST)
assert len(check.sections)==len(Document(SOURCE).sections)==6
assert len(check.inline_shapes)==6
mentions=[p.text for p in check.paragraphs if 'VALE3' in p.text]
assert len(mentions)==1 and 'continuidade' in mentions[0]
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
for old,new in zip(Document(SOURCE).inline_shapes,list(check.inline_shapes)[:4]):
    assert old.width==new.width and old.height==new.height
(OUT/'texto_acrescentado.md').write_text('\n\n'.join(texts),encoding='utf-8')
(OUT/'validacao_documento.json').write_text(json.dumps({'original_preservado':True,'original_sha256':source_hash,'figuras':6,'VALE3_somente_trabalhos_futuros':True,'documento':str(DEST),'metricas':m.to_dict()},ensure_ascii=False,indent=2),encoding='utf-8')
print(DEST)
