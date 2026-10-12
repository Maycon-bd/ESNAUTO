"""Extrai a Tabela 2 sem corrigir ou excluir estudos silenciosamente."""
import argparse
from collections import Counter
from pathlib import Path
import hashlib
import json
import re

from docx import Document


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("documento", type=Path)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    out = root / "reports" / "figures_bibliograficas"
    out.mkdir(parents=True, exist_ok=True)
    doc = Document(args.documento)
    tables = [t for t in doc.tables
              if len(t.columns) == 6 and t.cell(0, 0).text.strip() == "Trabalho"]
    if len(tables) != 1:
        raise ValueError("Não foi encontrada uma única matriz Trabalho × P1–P5")
    records = []
    for row in tables[0].rows[1:]:
        cells = [c.text.strip() for c in row.cells]
        years = re.findall(r"(?<!\d)(?:19|20)\d{2}(?!\d)", cells[0])
        if len(years) != 1 or any(c not in ("●", "○") for c in cells[1:]):
            raise ValueError(f"Linha exige conferência: {cells}")
        year = int(years[0])
        records.append({"estudo": cells[0], "ano_na_tabela": year,
                        "dimensoes": [int(c == "●") for c in cells[1:]],
                        "fora_recorte_2021_2026": not 2021 <= year <= 2026})
    counts = [sum(r["dimensoes"][i] for r in records) for i in range(5)]
    data = {"documento": args.documento.name,
            "sha256_documento": hashlib.sha256(args.documento.read_bytes()).hexdigest(),
            "origem": "Tabela 2 da monografia; anos e marcações transcritos sem alteração",
            "status": "Base descritiva da tabela atual, pendente de validação do portfólio",
            "n": len(records), "estudos": records, "cobertura": counts,
            "anos": dict(sorted(Counter(r["ano_na_tabela"] for r in records).items()))}
    (out / "base_tabela2.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"n": data["n"], "cobertura": counts, "anos": data["anos"]}))


if __name__ == "__main__":
    main()
