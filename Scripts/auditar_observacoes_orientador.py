from pathlib import Path
from zipfile import ZipFile
import re,json,hashlib
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'reports/observacoes_orientador_20261026'
record=json.loads((OUT/'registro_alteracoes.json').read_text(encoding='utf-8'))
source=Path(record['origem']);target=Path(record['saida'])
d=Document(target);src=Document(source)
notes={'Fonte estava diferente.','Tamanho da fonte estava 8,5 pts.','Tamanho da fonte da tabela...','Tamanho da fonte na tabela.'}
expected=[('' if p.text.strip() in notes else p.text) for p in src.paragraphs if p.text.strip() not in notes or p._p.xpath('.//w:sectPr')]
# A única mudança de texto posterior é a paginação estática das listas.
strip_page=lambda s:re.sub(r'\t\d+\s*$','\t#',s)
assert [strip_page(p.text) for p in d.paragraphs]==[strip_page(x) for x in expected]
assert [[c.text for row in t.rows for c in row.cells] for t in src.tables]==[[c.text for row in t.rows for c in row.cells] for t in d.tables]
paras=list(d.paragraphs)+[p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs]
bad=[];seen=0
for p in paras:
    chars=[(r.italic,r.font.subscript) for r in p.runs for c in r.text]
    text=''.join(r.text for r in p.runs)
    for m in re.finditer(r'(?<![\w])W(?:in|out)(?![\w])',text):
        seen+=1
        if any(not chars[k][0] for k in range(*m.span())) or any(not chars[k][1] for k in range(m.start()+1,m.end())):
            bad.append((m.group(),p.text[:120]))
assert not bad,bad
assert not [x for x in d._element.xpath('.//w:color') if x.get(qn('w:val'),'').upper()=='FF0000']
a=ZipFile(source);b=ZipFile(target)
imgs=[a.read(n) for n in a.namelist() if n.startswith('word/media/')]
imgs2=[b.read(n) for n in b.namelist() if n.startswith('word/media/')]
assert all(x in imgs2 for x in imgs)
assert hashlib.sha256(source.read_bytes()).hexdigest()==record['sha256_origem']
pdf=PdfReader(OUT/'monografia_observacoes_atendidas.pdf')
empty=[i+1 for i,p in enumerate(pdf.pages) if len(p.extract_text().strip())<15]
assert not empty,empty
audit={'conteudo_preservado_exceto_recados_e_paginas_listas':True,'figuras_preservadas':6,'tabelas_preservadas':13,
       'secoes_preservadas':len(d.sections),'variaveis_compostas_conferidas':seen,'sem_vermelho_remanescente':True,
       'paginas_pdf':len(pdf.pages),'paginas_vazias':empty,'origem_preservada':True,
       'sha256_docx_final':hashlib.sha256(target.read_bytes()).hexdigest()}
(OUT/'validacao_final.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False))
# Folhas de contato para inspeção da paginação integral; detalhes nas páginas PNG e pares.
from PIL import Image,ImageDraw
pages=[OUT/'final/png'/f'page-{i}.png' for i in range(1,len(pdf.pages)+1)]
for start in range(0,len(pages),12):
    sheet=Image.new('RGB',(1500,2920),'#d9dde0');draw=ImageDraw.Draw(sheet)
    for j,path in enumerate(pages[start:start+12]):
        im=Image.open(path);im.thumbnail((485,690));xx=(j%3)*500;yy=(j//3)*730
        sheet.paste(im,(xx,yy+25));draw.text((xx+5,yy+3),str(start+j+1),fill='black')
    sheet.save(OUT/'final'/f'contato-{start//12+1}.jpg')
