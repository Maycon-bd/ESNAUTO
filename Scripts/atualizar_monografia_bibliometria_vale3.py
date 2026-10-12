"""Insere a caracterização bibliográfica e a extensão VALE3 em uma cópia do DOCX."""
from copy import deepcopy
from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r'C:/Users/User/Downloads/MayconGarciaSilva_monografia_consolidada (1).docx')
OUTPUT = ROOT / 'reports/monografia/MayconGarciaSilva_monografia_bibliometria_VALE3.docx'
QA = ROOT / 'reports/qa_monografia_bibliometria'
FIG = QA / 'figuras'
d = Document(SOURCE)
original = list(d.paragraphs)
original_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()

def replace(p, text):
    rpr = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if rpr is not None: r._r.insert(0, rpr)

def before(anchor, text='', style='Normal', template=None):
    el = OxmlElement('w:p')
    anchor._p.addprevious(el)
    p = Paragraph(el, anchor._parent)
    p.style = style
    if template is not None and template._p.pPr is not None:
        p._p.insert(0,deepcopy(template._p.pPr))
    if text: p.add_run(text)
    return p

def prose(anchor,text):
    p=before(anchor,text)
    p.paragraph_format.widow_control=True
    for r in p.runs: r.font.name='Times New Roman'; r.font.size=Pt(12)
    return p

captions = [
 'Figura 1 – Cobertura das dimensões P1–P5 nos estudos',
 'Figura 2 – Frequência das palavras-chave declaradas pelos autores',
 'Figura 3 – Rede de coocorrência de palavras-chave',
 'Figura 4 – Expressões recorrentes nos títulos dos estudos',
]
# Lista de figuras, seguindo as listas preliminares já existentes.
h=before(original[171],'LISTA DE FIGURAS',template=original[171])
h.paragraph_format.page_break_before=True
h.paragraph_format.first_line_indent=Cm(0)
h.runs[0].bold=True
before(original[171])
for caption in captions:
    p=before(original[171],caption+'\t0',template=original[174])
    p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.line_spacing=1
    p.paragraph_format.space_after=Pt(12)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
original[171].paragraph_format.page_break_before=True
before(original[209],'VALE3\t - Ação ordinária da Vale S.A.',template=original[200])
# A lista de figuras acrescenta uma página preliminar. Manter a contagem
# institucional que omite a capa, ajustando o reinício da seção textual.
for el in d.sections[-1]._sectPr.findall(qn('w:pgNumType')):
    if el.get(qn('w:start')) is not None:
        el.set(qn('w:start'),str(int(el.get(qn('w:start')))+1))

replace(original[350], 'Como apoio à fundamentação teórica, foi conduzida uma extração estruturada de 34 estudos utilizados no trabalho. Essa coleção reúne estudos recentes considerados no processo de revisão e referências complementares de caráter conceitual, metodológico ou histórico, conforme a distinção descrita na Seção 3.8.3. Inclui artigos, trabalhos em eventos, preprints e trabalhos acadêmicos; por isso, seu quantitativo não deve ser interpretado como o número de artigos que satisfizeram todos os critérios de inclusão da busca recente. Lampugnani (2026) é discutido separadamente e não integra os 34 registros desta matriz. Cada estudo foi avaliado segundo cinco perguntas analíticas: P1 distribuição de probabilidade adotada na inicialização dos pesos do reservatório; P2 método de otimização dos hiperparâmetros da ESN; P3 etapas de pré-processamento das séries temporais; P4 comparação experimental com modelos de Deep Learning recorrentes (LSTM, GRU, Transformer); e P5 mercado financeiro e ativo investigado. A Tabela 1 sintetiza a cobertura dessas dimensões e os principais achados temáticos.')
replace(original[353], original[353].text.replace('cada artigo do portfólio','cada estudo da coleção').replace('dimens ões','dimensões').replace('que o artigo fornece','que o estudo fornece'))
replace(original[357], original[357].text.replace('34 artigos','34 estudos'))
replace(original[355],original[355].text.replace('artigo ×','estudo ×'))
for p in original[:255]:
    if p.text.startswith('Tabela 2 – Mapa de cobertura: artigo ×'):
        replace(p,p.text.replace('artigo ×','estudo ×'))

a=original[358]
prose(a,'A caracterização da coleção foi complementada pelas Figuras 1 a 4, elaboradas a partir da matriz de extração e de metadados reconstruídos dos documentos. Os formatos de frequência e a análise de termos dos títulos tomaram como referência as revisões de Cardoso et al. (2022c, 2024), utilizando os registros desta monografia. O processamento das matrizes foi realizado em R com bibliometrix, e a base foi validada para carregamento no Biblioshiny, sua interface gráfica (ARIA; CUCCURULLO, 2017; ARIA et al., 2026). O procedimento é detalhado na Seção 3.8.3.')

def figure(n,source,note,interpretation):
    p=before(a,captions[n-1],'Caption')
    pf=p.paragraph_format
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0)
    pf.line_spacing=1; pf.space_before=Pt(12); pf.space_after=Pt(6); pf.keep_with_next=True
    for r in p.runs: r.font.name='Times New Roman';r.font.size=Pt(10)
    p=before(a)
    pf=p.paragraph_format
    pf.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf.first_line_indent=Cm(0);pf.line_spacing=1;pf.keep_with_next=True
    p.add_run().add_picture(str(FIG / f'{n:02}.png'),width=Cm(16))
    inline=p.runs[0]._r.xpath('.//wp:docPr')[0]
    inline.set('descr',captions[n-1]+'. '+note)
    for label,text in [('Fonte: ',source),('Nota: ',note)]:
        p=before(a,label+text)
        pf=p.paragraph_format
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;pf.first_line_indent=Cm(0);pf.line_spacing=1
        pf.space_after=Pt(4);pf.keep_with_next=(label=='Fonte: ');pf.widow_control=True
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
    prose(a,interpretation)

figure(1,'Elaborado pelo autor a partir da Tabela 2 (2026).',
 'P1: inicialização dos pesos; P2: otimização de hiperparâmetros; P3: pré-processamento; P4: comparação com Deep Learning; P5: mercado e ativo. Um estudo pode cobrir várias dimensões. O denominador dos percentuais é 34; os valores não somam 100%.',
 'A Figura 1 sintetiza a cobertura das dimensões nos 34 estudos. O pré-processamento aparece em 31 estudos (91,2%), seguido de mercado e ativo, em 30 (88,2%). A comparação com Deep Learning é registrada em 21 estudos (61,8%), enquanto inicialização dos pesos e otimização dos hiperparâmetros são abordadas em 16 estudos cada (47,1%). Essas frequências permitem comparar a cobertura relativa das perguntas de extração. A identificação das distribuições efetivamente utilizadas e de sua fundamentação depende da leitura dos estudos e da matriz, além da contagem agregada.')
figure(2,'Elaborado pelo autor a partir dos PDFs citados, com processamento no bibliometrix 5.5.0 (2026).',
 'Foram analisados 24 documentos com palavras-chave declaradas localizadas, correspondentes a 70,6% dos 34 registros. São exibidos termos presentes em pelo menos dois documentos, com uma contagem por termo e documento.',
 'Na Figura 2, echo state network foi identificado em nove documentos (37,5% dos 24), e reservoir computing, em cinco (20,8%). Deep learning, stock price forecasting e time series aparecem em três documentos cada (12,5%). Os resultados caracterizam os temas declarados pelos autores no subconjunto com esse campo disponível. A ausência de um termo entre as palavras-chave não implica ausência de discussão do conceito no texto do estudo.')
figure(3,'Elaborado pelo autor com bibliometrix 5.5.0 a partir da base reconstruída (2026).',
 'Os mesmos 24 documentos da Figura 2 originam a matriz. A frequência mínima de dois documentos resulta em 14 termos e 33 ligações. A área dos nós representa frequência documental; a espessura das ligações, documentos compartilhados; as cores, comunidades Walktrap. A rede não representa coautoria ou citações.',
 'A Figura 3 mostra relações entre echo state network, reservoir computing, deep learning e termos associados à previsão financeira. LSTM e GRU também integram a rede, juntamente com optimization e machine learning. As ligações indicam declaração conjunta das palavras-chave nos documentos analisados. Essa caracterização temática complementa a leitura dos estudos; as afirmações específicas sobre inicialização de Win/W, otimização e pré-processamento continuam apoiadas na extração P1–P5.')
figure(4,'Elaborado pelo autor a partir dos 34 títulos da base reconstruída (2026).',
 'Cada par de palavras adjacentes é contado uma vez por documento. São exibidas 19 expressões presentes em pelo menos dois títulos. Singular, plural e siglas são conservados separadamente. O título inglês de Costa e Lunkes foi transcrito do próprio PDF. As expressões podem compartilhar palavras e ocorrer no mesmo documento.',
 'A Figura 4 apresenta 19 expressões recorrentes nos títulos. Echo state ocorre em 11 títulos (32,4%), state network em sete (20,6%) e time series em seis (17,6%). Reservoir computing, stock market e stock price aparecem em quatro títulos cada (11,8%), enquanto genetic algorithm ocorre em dois (5,9%). Echo state e state network são partes da denominação Echo State Network; suas frequências não devem ser somadas para estimar a quantidade de estudos sobre ESN.')
prose(a,'A análise dos títulos abrange 34 documentos, enquanto a análise de palavras-chave abrange 24. As diferenças entre suas frequências refletem também a cobertura dos campos, a redação dos títulos e as regras de normalização. Os gráficos descrevem a coleção utilizada no trabalho e não estimam a produção de toda a área nem demonstram, isoladamente, uma lacuna científica.')

# Extensão de dados: relato da implementação, sem inventar resultados.
a=original[370]
prose(a,'Como extensão funcional do ESNAUTO, foi incorporada a seleção entre PETR4 e VALE3, ação ordinária da Vale S.A. O arquivo fornecido para VALE3, identificado como série de fechamento com Fator de Ajuste de 2000–2020, contém 5.192 valores. Na configuração padrão, sua ordem é preservada e a divisão 50%/25%/25% corresponde a 2.596 observações de treino, 1.298 de validação e 1.298 de teste. Como o arquivo contém valores sem uma coluna de datas, o intervalo indicado no nome não permite atribuir datas exatas às partições nem demonstrar alinhamento diário com a PETR4. O procedimento de ajuste deve ser conferido na origem da série antes de uma comparação experimental entre ativos.')
prose(a,'A seleção da ação foi integrada à interface Shiny e à execução em lote, compartilhando a série e as partições entre ESN, otimização genética, LSTM e GRU. A camada de saída da ESN é ajustada aos dados selecionados. Até a consolidação desta versão, foram realizadas verificações funcionais de carregamento, troca do ativo e isolamento de resultados, além de uma execução reduzida da ESN; não foi executado o protocolo experimental completo na VALE3. Assim, as métricas, tabelas de desempenho e conclusões experimentais apresentadas nesta monografia continuam referentes à PETR4.')

replace(original[419],original[419].text.replace('A constituição do portfólio bibliográfico apresentado no Capítulo 2 seguiu','A seleção de estudos recentes para a revisão bibliográfica seguiu'))
a=original[423]
prose(a,'O recorte de 2021–2026 foi aplicado à seleção de estudos recentes. Referências anteriores a 2021 foram mantidas como fontes complementares quando necessárias à definição dos conceitos, à apresentação dos métodos originais ou à contextualização técnica e histórica. Sua utilização na fundamentação não significa que atenderam ao critério temporal da busca recente. A coleção descritiva de 34 estudos das Tabelas 1 e 2 e das Figuras 1 a 4 incorpora essas fontes complementares, distinguindo-se do conjunto estritamente elegível pelos critérios CI1–CI5.')
prose(a,'Para a análise bibliográfica complementar, os 34 registros foram reconstruídos a partir dos PDFs disponíveis e das referências da monografia, com conferência de metadados por DOI ou título. Os registros foram identificados como uma coleção reconstruída, sem atribuir a ela a condição de exportação original do Scopus. Foram preservadas as informações das versões locais consultadas e as ressalvas sobre preprints ou outras versões. Contagens de citações recebidas e listas completas de referências citadas não foram reconstruídas; por esse motivo, não foram produzidos rankings de impacto ou redes de cocitação.')
prose(a,'O processamento utilizou bibliometrix 5.5.0 em R, pacote para análise bibliométrica e mapeamento científico descrito por Aria e Cuccurullo (2017). A coleção foi salva em formato compatível com Biblioshiny e validada com o carregador dessa interface, apresentada por Aria et al. (2026). As matrizes foram calculadas em scripts de R, com a função biblioNetwork, e a apresentação das figuras utilizou ggplot2 e igraph. A validação de carregamento não corresponde a uma execução das análises pela interface gráfica do Biblioshiny. As revisões de Cardoso et al. (2022c, 2024) serviram como referências para a escolha dos formatos de visualização.')
prose(a,'A frequência de palavras-chave considerou os 24 documentos em que foram localizadas listas declaradas pelos autores. Variantes equivalentes foram normalizadas com registro das correspondências, sem converter títulos ou classificações P1–P5 em palavras-chave. Cada termo contou uma vez por documento. A matriz de coocorrência foi conferida por uma contagem binária independente; o corte de frequência mínima de dois documentos resultou em 14 nós e 33 ligações. O posicionamento visual utilizou Fruchterman–Reingold, e as comunidades foram identificadas por Walktrap, com semente 20261012.')
prose(a,'Para os títulos, foram extraídos pares de palavras adjacentes (bigrams) dos 34 registros. O texto foi convertido para minúsculas e tokenizado, sem conservar pontuação, algarismos ou hífens nos tokens. Os pares foram formados antes da exclusão de palavras de ligação, evitando unir palavras originalmente distantes. Não foram aplicados stemming, expansão de siglas nem agrupamento entre singular e plural. Cada par contou uma vez por documento, com frequência mínima de dois títulos para exibição. Para Costa e Lunkes, utilizou-se o título inglês apresentado na página 2 do PDF. A matriz foi processada com biblioNetwork e conferida por contagem independente. Os títulos originais, regras de tratamento e identificadores foram preservados para rastreabilidade.')

a=original[489]
prose(a,'A inclusão da VALE3 ampliou o software com um seletor explícito da ação utilizada em cada execução. A troca de ativo invalida resultados exibidos da série anterior e é bloqueada durante uma execução em andamento. Históricos e recordes são separados por ação, e a execução em lote registra a identificação do dataset, seu hash e os tamanhos das partições. Esses mecanismos reduzem o risco de comparar ou apresentar resultados de ativos diferentes como se pertencessem ao mesmo experimento. As verificações funcionais dessa extensão demonstram o funcionamento do fluxo de seleção, mas não estabelecem desempenho preditivo na VALE3.')
replace(original[500],original[500].text.replace('A pesquisa considera um único ativo','A avaliação experimental considera um único ativo'))
replace(original[502],'Em relação ao software, o ESNAUTO passou a disponibilizar a seleção entre PETR4 e VALE3. Entretanto, a validação científica apresentada nesta monografia permanece concentrada na PETR4, pois a VALE3 ainda não foi submetida ao protocolo experimental completo. A ampliação funcional permite organizar novas execuções, mas não sustenta conclusões sobre generalização entre ativos antes de sua avaliação.')
prose(original[508],'Como próxima etapa experimental, propõe-se executar na VALE3 o protocolo de ESN, otimização genética, LSTM e GRU, preservando previsões, métricas, sementes, matrizes e metadados por ação. A comparação deverá explicitar os tamanhos e a correspondência temporal das partições, além de empregar procedimentos de pré-processamento ajustados apenas ao treino. Os parâmetros de distribuições ajustados à PETR4, inclusive os utilizados na Onda 5, não devem ser assumidos como estimados para a VALE3: sua aplicação requer nova estimação no treino da VALE3 ou identificação explícita de um experimento de transferência. Essa avaliação permitirá acrescentar os gráficos de previsão e discutir os resultados do segundo ativo em uma atualização posterior.')
replace(original[513],original[513].text+' A extensão do seletor para PETR4 e VALE3 acrescentou reutilização do fluxo e separação dos resultados por ação, permanecendo pendente a avaliação experimental completa da VALE3.')

# Referências citadas nas novas passagens. Sufixos históricos 2022a/2022b preservados.
refs_aria=[
 'ARIA, M.; CUCCURULLO, C. bibliometrix: an R-tool for comprehensive science mapping analysis. Journal of Informetrics, v. 11, n. 4, p. 959-975, 2017. DOI: 10.1016/j.joi.2017.08.007.',
 "ARIA, M.; CUCCURULLO, C.; D’ANIELLO, L.; SPANO, M. Biblioshiny and the SAAS Workflow: an integrated framework for transparent and reproducible science mapping. A demonstration through the replication of a study. Journal of Informetrics, v. 20, n. 3, art. 101837, 2026. DOI: 10.1016/j.joi.2026.101837.",
]
refs_cardoso=[
 'CARDOSO, F. C.; BERRI, R. A.; BORGES, E. N.; LUCCA, G.; MATTOS, V. L. D. Associação de redes de estado de eco com métodos estatísticos clássicos de previsão: uma revisão bibliométrica. In: XXV ENCONTRO NACIONAL DE MODELAGEM COMPUTACIONAL, 2022, evento virtual. Anais [...]. 2022c.',
 'CARDOSO, F. C.; BERRI, R. A.; BORGES, E. N.; DALMAZO, B. L.; LUCCA, G.; MATTOS, V. L. D. Echo state network and classical statistical techniques for time series forecasting: a review. Knowledge-Based Systems, v. 293, art. 111639, 2024. DOI: 10.1016/j.knosys.2024.111639.',
]
for anchor,refs in [(original[518],refs_aria),(original[521],refs_cardoso)]:
    for ref in refs:
        p=before(anchor,ref,original[517].style.name,original[517])
        p.paragraph_format.keep_together=True
        p.paragraph_format.widow_control=True
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
d.save(OUTPUT)
check=Document(OUTPUT)
assert len(check.tables)==13 and len(check.inline_shapes)==4
assert len(check.sections)==6
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==original_hash
for s,t in zip(Document(SOURCE).sections,check.sections):
    assert all(getattr(s,k)==getattr(t,k) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin'])
QA.mkdir(parents=True,exist_ok=True)
(QA/'validacao_edicao.json').write_text(json.dumps({'original_sha256':original_hash,'original_preservado':True,'tabelas':len(check.tables),'figuras':len(check.inline_shapes),'secoes':len(check.sections),'output':str(OUTPUT)},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUTPUT)
