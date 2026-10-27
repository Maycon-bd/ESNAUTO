"""Revisão tipográfica do Word devolvido pelo orientador, preservando conteúdo e evidências."""
from pathlib import Path
from copy import deepcopy
import re,json,hashlib
from collections import Counter
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('C:/Users/User/Downloads/MayconGarciaSilva_monografia_revisao_orientador com observações.docx')
DEST=ROOT/'reports/monografia/MayconGarciaSilva_monografia_observacoes_atendidas.docx'
OUT=ROOT/'reports/observacoes_orientador_20261026'
OUT.mkdir(parents=True,exist_ok=True)
source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
d=Document(SOURCE)
original=list(d.paragraphs)
notes=['Fonte estava diferente.','Tamanho da fonte estava 8,5 pts.','Tamanho da fonte da tabela...','Tamanho da fonte na tabela.']
changes=[]
for p in original:
    if p.text.strip() in notes:
        changes.append({'tipo':'recado_atendido','antes':p.text,'depois':'Removido da cópia limpa; atendimento de fonte e tamanho registrado.'})
        if p._p.xpath('.//w:sectPr'):
            p.clear()  # Preservar a quebra de seção existente junto ao recado.
        else:
            p._p.getparent().remove(p._p)

# Os trechos exemplificados pelo orientador indicam W em itálico e in/out subscritos, também em itálico.
# Divide somente runs de texto simples; preserva campos, equações OMML, desenhos e propriedades existentes.
foreign=[
 'Cross-generational elitist selection, Heterogeneous recombination and Cataclysmic mutation',
 'Backpropagation Through Time','Latin Hypercube Sampling','Half Uniform Crossover',
 'Normalized Root Mean Squared Error','Normalized Mean Absolute Error','Mean Absolute Percentage Error',
 'Mean Absolute Error','Root Mean Squared Error','Generalized Error Distribution',
 'Higher Order Neural Networks','Higher-Order Neural Networks','Recurrent Neural Networks','Recurrent Neural Network',
 'Long-Short Term Memory','Long Short-Term Memory','Long Short Term Memory','Long-Short-Term Memory',
 'Gated Recurrent Units','Gated Recurrent Unit','Support Vector Machine','Support Vector Machines',
 'Large Language Models','Multilayer Perceptron','Particle Swarm Optimization',
 'Artificial Neural Networks','Echo State Networks','Echo State Network','Echo State Property',
 'reservoir computing','reservoir pool','deep learning','machine learning','genetic algorithm',
 'stock market prediction','stock price forecasting','financial market prediction','financial time series',
 'time series','state network','echo state','stock market','stock price',
 'Power Function','leaking rate','leak rate','spectral radius','washout','fitness',
 'baselines','baseline','benchmarks','benchmark','datasets','dataset','scripts','script',
 'preprints','preprint','software','hardware','dropout','backpropagation','crossover',
 'overfitting','underfitting','offline','online','rollback','pipeline','framework',
 'backtesting','feedback','feedforward','out-of-sample','in-sample','optimization',
 'training','validation','forecasting','forecast','pool','seed','seeds','bigram','bigrams',
 'stemming','stemming','stopwords','bias','input','output','grid search','random search',
 'learning rate','batch size','early stopping','holdout','off-line','on-line',
 'Sparrow Search Algorithm','Transformer','Transformers','Min-Max','z-score','data leakage'
]
pattern=re.compile(r'(?<![\w])(?:'+ '|'.join(re.escape(x) for x in sorted(set(foreign),key=len,reverse=True))+r')(?![\w])',re.I)
varpattern=re.compile(r'(?<![\w])(?:Wout|Win|W)(?![\w])')
count=Counter()

def format_paragraph(p,do_foreign=True,do_vars=True):
    runs=p._p.xpath('.//w:r[not(ancestor::w:del)]')
    # w:t únicos e runs puramente textuais: não modificar instruções de campo nem desenhos.
    items=[];offset=0
    for r in runs:
        ts=r.findall(qn('w:t'))
        txt=''.join(x.text or '' for x in ts)
        safe=len(ts)==1 and all(x.tag in {qn('w:rPr'),qn('w:t')} for x in r)
        items.append((r,offset,offset+len(txt),txt,safe));offset+=len(txt)
    text=''.join(x[3] for x in items)
    spans=[]
    if do_foreign:
        for m in pattern.finditer(text):spans.append((m.start(),m.end(),'foreign'));count['estrangeirismos']+=1
    for m in (varpattern.finditer(text) if do_vars else []):
        if m.group()=='W' and re.search(r'\bWang\s+$',text[:m.start()]):continue
        a,z=m.span();spans.append((a,a+1,'W'))
        if z>a+1:spans.append((a+1,z,'sub'))
        count['variaveis']+=1
    for r,a,z,txt,safe in items:
        if not safe:continue
        cuts={0,len(txt)}
        intersect=[(x,y,k) for x,y,k in spans if x<z and y>a]
        if not intersect:continue
        for x,y,k in intersect:cuts.add(max(0,x-a));cuts.add(min(len(txt),y-a))
        cuts=sorted(cuts)
        for lo,hi in zip(cuts,cuts[1:]):
            if lo==hi:continue
            clone=deepcopy(r);t=clone.find(qn('w:t'));t.text=txt[lo:hi]
            t.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
            kinds={k for x,y,k in intersect if x<a+hi and y>a+lo}
            if kinds:
                pr=clone.find(qn('w:rPr'))
                if pr is None:pr=OxmlElement('w:rPr');clone.insert(0,pr)
                for name in ['i','iCs']:
                    node=pr.find(qn('w:'+name))
                    if node is None:node=OxmlElement('w:'+name);pr.append(node)
                    node.set(qn('w:val'),'1')
                if 'sub' in kinds or 'W' in kinds:
                    node=pr.find(qn('w:vertAlign'))
                    if node is None:node=OxmlElement('w:vertAlign');pr.append(node)
                    node.set(qn('w:val'),'subscript' if 'sub' in kinds else 'baseline')
            r.addprevious(clone)
        r.getparent().remove(r)

refs=False;abstract=False
for p in d.paragraphs:
    if p.text.strip()=='ABSTRACT':abstract=True
    if abstract and p.text.strip()=='LISTA DE TABELAS':abstract=False
    if p.text.strip()=='REFERÊNCIAS' and '\t' not in p.text:refs=True
    format_paragraph(p,not refs and not abstract,not refs)
    # O sumário foi corrigido para a fonte do corpo pelo orientador.
    if '\t' in p.text and re.search(r'\t\d+\s*$',p.text):
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
        if re.match(r'^\d+(?:\.\d+)*\t',p.text):
            p.paragraph_format.left_indent=Cm(1.2)
            p.paragraph_format.first_line_indent=Cm(-1.2)
            p.paragraph_format.right_indent=Cm(.75)
            tabs=p.paragraph_format.tab_stops;tabs.clear_all()
            tabs.add_tab_stop(Cm(1.2))
            tabs.add_tab_stop(Cm(15),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)

# Os estilos no Word devolvido tinham numId=0, ocultando os indicativos que constam do sumário.
# Restaurar a numeração multinível já existente (numId=1), sem digitar números no conteúdo.
for sty,level in [('TITULO TCC',0),('SEÇÃO TCC',1),('SUBÇÃO 2',2)]:
    pr=d.styles[sty]._element.get_or_add_pPr()
    num=pr.find(qn('w:numPr'))
    if num is None:num=OxmlElement('w:numPr');pr.append(num)
    for tag,value in [('ilvl',level),('numId',1)]:
        node=num.find(qn('w:'+tag))
        if node is None:node=OxmlElement('w:'+tag);num.append(node)
        node.set(qn('w:val'),str(value))
for p in d.paragraphs:
    if p.text=='REFERÊNCIAS':
        num=p._p.get_or_add_pPr().find(qn('w:numPr'))
        if num is None:num=OxmlElement('w:numPr');p._p.pPr.append(num)
        node=OxmlElement('w:numId');node.set(qn('w:val'),'0');num.append(node)
changes.append({'tipo':'numeracao_secoes','motivo':'Restaurar indicativos automáticos dos capítulos, seções e subseções, ocultos nos estilos da origem, de acordo com o sumário.'})

for ti,table in enumerate(d.tables):
    if ti==0:continue  # Ficha catalográfica vazia.
    size=12 if ti in [1,2] else 11
    for ri,row in enumerate(table.rows):
        # Repetir o cabeçalho quando a tabela for distribuída em mais de uma página.
        pr=row._tr.get_or_add_trPr()
        if ri==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in c.paragraphs:
                format_paragraph(p)
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.widow_control=True
                # Colunas textuais longas à esquerda, colunas numéricas centralizadas.
                descriptive=(ti==1 and ci in [0,2]) or (ti==2 and ci==0) or (ti in [9,11,12])
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if descriptive else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(size)
    changes.append({'tipo':'fonte_tabela','tabela_ooxml':ti,'tamanho_final_pt':size,'motivo':'Preservar 12 pt nas Tabelas 1 e 2 corrigidas pelo orientador; ampliar tabelas antes em 6,5–10 pt para 11 pt.'})

# Trocar apenas o vermelho de revisão por preto; outras propriedades permanecem.
for color in d._element.xpath('.//w:color'):
    val=color.get(qn('w:val'),'').upper()
    if val=='FF0000':color.set(qn('w:val'),'000000');count['marcacoes_vermelhas_limpas']+=1

for name in ['Normal','Caption']:
    d.styles[name].font.name='Times New Roman'
d.save(DEST)
check=Document(DEST)
expected=[('' if p.text.strip() in notes else p.text) for p in Document(SOURCE).paragraphs if p.text.strip() not in notes or p._p.xpath('.//w:sectPr')]
assert [p.text for p in check.paragraphs]==expected
assert len(check.tables)==len(Document(SOURCE).tables)==13
assert len(check.inline_shapes)==len(Document(SOURCE).inline_shapes)==6
assert len(check.sections)==len(Document(SOURCE).sections)==6
assert [[c.text for row in t.rows for c in row.cells] for t in check.tables]==[[c.text for row in t.rows for c in row.cells] for t in Document(SOURCE).tables]
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
record={'origem':str(SOURCE),'sha256_origem':source_hash,'saida':str(DEST),'alteracoes':changes,'contagens':dict(count),
        'conteudo_preservado_exceto_recados':True,'imagens':6,'tabelas':13,'secoes':6,
        'observacoes':'Base explícita devolvida pelo orientador em 26/10/2026; snapshot antigo não foi renovado. Alterações editoriais sem reexecução de experimentos.'}
(OUT/'registro_alteracoes.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'registro_alteracoes.md').write_text('# Atendimento às observações do orientador\n\nOrigem: '+str(SOURCE)+'\n\nSHA-256: '+source_hash+'\n\n- Preservadas as correções de texto e notação já feitas em vermelho.\n- Padronizados estrangeirismos no texto em português e variáveis W, Win e Wout (W em itálico, in/out em subscrito itálico).\n- Mantidos 12 pt nas duas primeiras tabelas; demais tabelas ampliadas para 11 pt, com cabeçalhos repetidos.\n- Removidos da cópia limpa somente os quatro recados editoriais identificados.\n- Conteúdo, números, tabelas e seis figuras preservados; não houve nova execução de modelos.\n- O sumário será atualizado pela paginação final do Word.\n',encoding='utf-8')
print(DEST);print(dict(count))
