from docx import Document
from pathlib import Path
import json
p=Path(r'C:/Users/User/Downloads/MayconGarciaSilva_monografia_consolidada (1).docx')
d=Document(p)
for i, q in enumerate(d.paragraphs):
 if i>250 and (q.style.name!='Normal' or any(x in q.text for x in ['VALE','portfólio','REPRODU','CONTRIBUI','TRABALHOS FUTUROS','único ativo','Bibliometrix'])):
  print(i,q.style.name,q.text[:180])
print('sections',len(d.sections),'tables',len(d.tables),'pictures',len(d.inline_shapes))
for s in d.sections:
 print('Page/margins', [round(x.cm,2) for x in [s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin]])
for i in [350,355,358,419,484,489,500,502,505,513,515,520]:
 q=d.paragraphs[i]; print('STYLEDETAIL',i,q.style.name,'ppr',q._p.pPr.xml if q._p.pPr is not None else None,'run',[r._r.rPr.xml if r._r.rPr is not None else None for r in q.runs[:1]])
